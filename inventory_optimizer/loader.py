"""Carga del Excel de entrada, tolerante a columnas faltantes o con otro nombre.

Todo lo que falta se completa con los defaults de `Params` y queda
registrado en `Dataset.warnings`, que el reporte muestra en una hoja propia.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import fields
from typing import Any, BinaryIO

from openpyxl import load_workbook

from .model import Customer, Dataset, Lane, Params, Plant, Process, Product, StockRecord
from .schema import PARAMS_SHEET_ALIASES, SHEETS, Sheet, norm

_MONTHS_ES = {
    "ene": 1, "jan": 1, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "ago": 8, "aug": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12, "dec": 12,
}


# ---------------------------------------------------------------- parsing

def to_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().replace(" ", "")
    pct = s.endswith("%")
    s = s.rstrip("%")
    if "," in s and "." in s:          # 1.234,5 o 1,234.5
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return None
    return x / 100 if pct else x


def to_fraction(value: Any) -> float | None:
    """Acepta 57, 57%, 0.57 -> 0.57."""
    x = to_float(value)
    if x is None:
        return None
    return x / 100 if x > 1 else x


def to_month(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, (dt.date, dt.datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    s = norm(value)
    m = re.fullmatch(r"(\d{4})_(\d{1,2})(?:_\d{1,2})?", s)
    if m:
        return f"{int(m[1]):04d}-{int(m[2]):02d}"
    m = re.fullmatch(r"(?:\d{1,2}_)?(\d{1,2})_(\d{4})", s)
    if m:
        return f"{int(m[2]):04d}-{int(m[1]):02d}"
    m = re.fullmatch(r"([a-z]+)_?(\d{2,4})", s)
    if m and m[1][:3] in _MONTHS_ES:
        y = int(m[2])
        y = y + 2000 if y < 100 else y
        return f"{y:04d}-{_MONTHS_ES[m[1][:3]]:02d}"
    m = re.fullmatch(r"(\d{4})(\d{2})", s)
    if m and 1 <= int(m[2]) <= 12:
        return f"{m[1]}-{m[2]}"
    return None


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


# ---------------------------------------------------------------- lectura de hojas

def _find_sheet(wb, aliases: tuple[str, ...]):
    by_norm = {norm(n): n for n in wb.sheetnames}
    for a in aliases:
        if a in by_norm:
            return wb[by_norm[a]]
    for n_norm, n in by_norm.items():       # coincidencia parcial: "Demanda 2024-2025"
        if any(n_norm.startswith(a) for a in aliases):
            return wb[n]
    return None


def _rows(ws) -> tuple[list[str], list[tuple]]:
    """Encabezado = primera fila con >=2 celdas de texto dentro de las 10 primeras."""
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows[:10]):
        texts = [c for c in row if isinstance(c, str) and c.strip()]
        if len(texts) >= 2:
            header = [norm(c) if c is not None else "" for c in row]
            body = [r for r in rows[i + 1:] if any(c not in (None, "") for c in r)]
            return header, body
    return [], []


def _map_columns(sheet: Sheet, header: list[str]) -> dict[str, int]:
    out: dict[str, int] = {}
    used: set[int] = set()
    for col in sheet.cols:               # primero nombres exactos
        if col.key in header:
            out[col.key] = header.index(col.key)
            used.add(out[col.key])
    for col in sheet.cols:               # luego sinónimos
        if col.key in out:
            continue
        for alias in col.aliases:
            if alias in header and header.index(alias) not in used:
                out[col.key] = header.index(alias)
                used.add(out[col.key])
                break
    return out


def _records(wb, sheet: Sheet, ds: Dataset) -> list[dict[str, Any]] | None:
    ws = _find_sheet(wb, sheet.aliases)
    if ws is None:
        return None
    header, body = _rows(ws)
    cmap = _map_columns(sheet, header)
    missing = [c.key for c in sheet.cols if c.required and c.key not in cmap]
    if sheet.name == "Demanda" and missing and set(missing) <= {"mes", "cantidad_t"}:
        return _wide_demand(header, body, cmap, ds)
    if missing:
        ds.warn(f"Hoja '{ws.title}': faltan columnas obligatorias {missing}; hoja ignorada.")
        return None
    unknown = [h for i, h in enumerate(header) if h and i not in cmap.values()]
    if unknown:
        ds.warn(f"Hoja '{ws.title}': columnas no reconocidas (ignoradas): {', '.join(unknown)}")
    return [{k: (row[i] if i < len(row) else None) for k, i in cmap.items()} for row in body]


def _wide_demand(header, body, cmap, ds) -> list[dict[str, Any]] | None:
    month_cols = {i: to_month(h.replace("_", "-")) for i, h in enumerate(header)}
    month_cols = {i: m for i, m in month_cols.items() if m}
    if not month_cols or "cliente" not in cmap or "producto" not in cmap:
        ds.warn("Hoja 'Demanda': no se reconoció ni formato largo (mes, cantidad) ni ancho (un mes por columna).")
        return None
    out = []
    for row in body:
        for i, month in month_cols.items():
            out.append({"cliente": row[cmap["cliente"]], "producto": row[cmap["producto"]],
                        "planta": row[cmap["planta"]] if "planta" in cmap else None,
                        "mes": month, "cantidad_t": row[i] if i < len(row) else None})
    return out


# ---------------------------------------------------------------- canonicalización de nombres

class _Names:
    def __init__(self):
        self.by_norm: dict[str, str] = {}

    def add(self, name: str) -> str:
        return self.by_norm.setdefault(norm(name), name)

    def get(self, name: str) -> str:
        return self.by_norm.get(norm(name), name)


def _location(value: Any, product: str) -> str:
    s = norm(value)
    if "concent" in s or "mp" == s or "materia" in s or "concent" in norm(product):
        return "concentrado"
    if "trans" in s or "ruta" in s or "embarc" in s:
        return "transito"
    if "proc" in s or "wip" in s or "curso" in s:
        return "proceso"
    if "consig" in s or "cliente" in s or "hub" in s or "deposito" in s:
        return "consignacion"
    return "pt"


# ---------------------------------------------------------------- carga principal

def load_excel(source: str | BinaryIO) -> Dataset:
    wb = load_workbook(source, data_only=True, read_only=True)
    ds = Dataset()
    _load_params(wb, ds)
    sheets = {s.name: s for s in SHEETS}
    plants, products, customers = _Names(), _Names(), _Names()

    for r in _records(wb, sheets["Plantas"], ds) or []:
        name = _text(r.get("planta"))
        if not name:
            continue
        ds.plants[plants.add(name)] = Plant(
            name=plants.get(name), country=_text(r.get("pais")), region=_text(r.get("region")),
            capacity_t_month=to_float(r.get("capacidad_t_mes")),
            conc_lead_time_days=to_float(r.get("lead_time_concentrado_dias")),
            conc_lead_time_sd_days=to_float(r.get("desv_lead_time_concentrado_dias")),
            conc_receipt_interval_days=to_float(r.get("intervalo_recepcion_concentrado_dias")),
            conc_stock_t=to_float(r.get("stock_concentrado_t")),
            conc_cost_usd_t=to_float(r.get("costo_concentrado_usd_t")),
        )

    for r in _records(wb, sheets["Productos"], ds) or []:
        name = _text(r.get("producto"))
        if not name:
            continue
        ds.products[products.add(name)] = Product(
            name=products.get(name), mo_content=to_fraction(r.get("contenido_mo_pct")),
            conc_factor=to_float(r.get("factor_concentrado")), value_usd_t=to_float(r.get("valor_usd_t")),
        )

    for r in _records(wb, sheets["Proceso"], ds) or []:
        pl, pr = _text(r.get("planta")), _text(r.get("producto"))
        if not pl or not pr:
            continue
        pl, pr = _ensure_plant(ds, plants, pl), _ensure_product(ds, products, pr)
        ds.processes[(pl, pr)] = Process(
            plant=pl, product=pr, process_days=to_float(r.get("tiempo_proceso_dias")),
            process_sd_days=to_float(r.get("desv_proceso_dias")), qa_days=to_float(r.get("tiempo_qa_dias")),
            campaign_interval_days=to_float(r.get("intervalo_campana_dias")),
            capacity_t_month=to_float(r.get("capacidad_t_mes")), cost_usd_t=to_float(r.get("costo_conversion_usd_t")),
        )

    for r in _records(wb, sheets["Clientes"], ds) or []:
        name = _text(r.get("cliente"))
        if not name:
            continue
        plant = _text(r.get("planta_asignada"))
        ds.customers[customers.add(name)] = Customer(
            name=customers.get(name), country=_text(r.get("pais")), region=_text(r.get("region")),
            plant=_ensure_plant(ds, plants, plant) if plant else "",
            segment=_text(r.get("segmento")).upper()[:1], service_level=to_fraction(r.get("nivel_servicio")),
            dispatch_interval_days=to_float(r.get("intervalo_despacho_dias")),
            firm_order_days=to_float(r.get("pedido_firme_dias")),
        )

    demand_rows = _records(wb, sheets["Demanda"], ds)
    if demand_rows is None:
        ds.warn("No se encontró la hoja 'Demanda': sin historia de demanda no se puede calcular stock óptimo.")
    for r in demand_rows or []:
        cu, pr, month = _text(r.get("cliente")), _text(r.get("producto")), to_month(r.get("mes"))
        qty = to_float(r.get("cantidad_t"))
        if not cu or not pr or not month or qty is None:
            continue
        if norm(cu) not in customers.by_norm:
            ds.customers[customers.add(cu)] = Customer(name=cu)
            ds.warn(f"Cliente '{cu}' aparece en Demanda pero no en Clientes; se creó con datos por defecto.")
        cu, pr = customers.get(cu), _ensure_product(ds, products, pr)
        series = ds.demand.setdefault((cu, pr), {})
        series[month] = series.get(month, 0.0) + qty
        plant = _text(r.get("planta"))
        if plant:
            ds.line_plant[(cu, pr)] = _ensure_plant(ds, plants, plant)

    for r in _records(wb, sheets["Transito"], ds) or []:
        pl, dest, days = _text(r.get("planta")), _text(r.get("destino")), to_float(r.get("dias_transito"))
        if not pl or not dest or days is None:
            continue
        ds.lanes.append(Lane(plant=_ensure_plant(ds, plants, pl), destination=dest, transit_days=days,
                             transit_sd_days=to_float(r.get("desv_transito_dias")),
                             freight_usd_t=to_float(r.get("flete_usd_t"))))

    for r in _records(wb, sheets["Stock"], ds) or []:
        pl, pr, qty = _text(r.get("planta")), _text(r.get("producto")), to_float(r.get("cantidad_t"))
        if not pl or not pr or qty is None:
            continue
        loc = _location(r.get("ubicacion"), pr)
        pl = _ensure_plant(ds, plants, pl)
        pr = pr if loc == "concentrado" else _ensure_product(ds, products, pr)
        ds.stock.append(StockRecord(plant=pl, product=pr, location=loc, qty_t=qty,
                                    month=to_month(r.get("mes")) or ""))

    wb.close()
    _validate(ds)
    return ds


def _ensure_plant(ds: Dataset, names: _Names, name: str) -> str:
    if norm(name) not in names.by_norm:
        names.add(name)
        ds.plants[name] = Plant(name=name)
        ds.warn(f"Planta '{name}' no está en la hoja Plantas; se creó con datos por defecto.")
    return names.get(name)


def _ensure_product(ds: Dataset, names: _Names, name: str) -> str:
    if norm(name) not in names.by_norm:
        names.add(name)
        ds.products[name] = Product(name=name)
    return names.get(name)


def _load_params(wb, ds: Dataset) -> None:
    ws = _find_sheet(wb, PARAMS_SHEET_ALIASES)
    if ws is None:
        ds.warn("Sin hoja 'Parametros': se usan los supuestos por defecto (ver hoja Supuestos del reporte).")
        return
    valid = {f.name: f for f in fields(Params)}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or len(row) < 2:
            continue
        key = norm(row[0])
        if key not in valid or row[1] is None:
            continue
        if valid[key].type in ("bool", bool):
            setattr(ds.params, key, str(row[1]).strip().lower() in ("1", "true", "si", "sí", "yes", "verdadero"))
            continue
        val = to_fraction(row[1]) if key.startswith("nivel_servicio") or key in (
            "tasa_costo_capital", "contenido_mo_concentrado", "utilizacion_alta", "utilizacion_baja") else to_float(row[1])
        if val is not None:
            setattr(ds.params, key, val)


def _validate(ds: Dataset) -> None:
    if not ds.demand:
        ds.warn("No hay filas de demanda válidas (cliente, producto, mes, cantidad): no se puede calcular "
                "stock óptimo ni rotación.")
    for (cu, pr) in ds.demand:
        c = ds.customers[cu]
        plant = ds.line_plant.get((cu, pr)) or c.plant
        if not plant:
            ds.warn(f"Cliente '{cu}' sin planta asignada; se le asignará la planta de menor costo.")
    for p in ds.products.values():
        if p.mo_content is None and p.value_usd_t is None:
            ds.warn(f"Producto '{p.name}' sin contenido de Mo ni valor: se valoriza con 57% Mo.")
    if not ds.processes:
        ds.warn("Sin hoja 'Proceso': se asume que todas las plantas fabrican todos los productos "
                "con tiempos de proceso por defecto.")
    if not ds.lanes:
        ds.warn("Sin hoja 'Transito': se usan días de tránsito por defecto para todas las rutas.")
    if not ds.stock:
        ds.warn("Sin hoja 'Stock': no se puede calcular rotación actual, solo el stock óptimo.")
