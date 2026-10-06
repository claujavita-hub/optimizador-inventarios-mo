"""Reporte Excel del análisis de rotación y stock óptimo."""
from __future__ import annotations

import datetime as dt
from collections import defaultdict
from dataclasses import fields

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .allocation import AllocationResult
from .engine import Analysis
from .recommendations import Recommendation
from .schema import PARAM_HELP

HEAD = PatternFill("solid", fgColor="1F4E78")
STATUS_FILL = {
    "Sobrestock": "F8CBAD", "Sobre máximo": "FCE4D6", "Riesgo de quiebre": "FF9999",
    "En rango": "C6EFCE", "Stock sin demanda": "D9D9D9", "Sin dato": "EDEDED",
}
PRIO_FILL = {"Alta": "FF9999", "Media": "FFE699", "Baja": "DDEBF7"}
T, USD, PCT, X2, D0 = "#,##0", '#,##0" USD"', "0.0%", "0.00", "0"


def _table(ws, row: int, headers: list[str], rows: list[list], fmts: list[str | None] | None = None,
           status_col: int | None = None, fills: dict | None = None) -> int:
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32
    for i, values in enumerate(rows, start=row + 1):
        for j, v in enumerate(values, start=1):
            c = ws.cell(i, j, v)
            if fmts and j - 1 < len(fmts) and fmts[j - 1] and isinstance(v, (int, float)):
                c.number_format = fmts[j - 1]
        if status_col is not None:
            key = values[status_col]
            color = (fills or STATUS_FILL).get(key)
            if color:
                ws.cell(i, status_col + 1).fill = PatternFill("solid", fgColor=color)
    return row + len(rows) + 2


def _widths(ws, widths: dict[int, float] | None = None, default: float = 14) -> None:
    for j in range(1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, default)


def detail_tables(alloc: AllocationResult, recs: list[Recommendation]) -> dict[str, dict]:
    """Tablas de detalle (encabezados, filas, formatos) que comparten el Excel y la app."""
    a = alloc.base
    p = a.ds.params
    out: dict[str, dict] = {}

    out["Recomendaciones"] = dict(
        headers=["#", "Prioridad", "Tipo", "Ámbito", "Diagnóstico", "Acción propuesta", "Impacto (t)",
                 "Capital liberado (+) / requerido (-) USD", "Ahorro anual USD"],
        rows=[[i, r.prioridad, r.tipo, r.ambito, r.diagnostico, r.accion, r.impacto_t, r.capital_usd,
               r.ahorro_anual_usd] for i, r in enumerate(recs, start=1)],
        fmts=[None, None, None, None, None, None, T, USD, USD], status_col=1, fills=PRIO_FILL, wrap=(5, 6),
        widths={1: 4, 2: 10, 3: 24, 4: 30, 5: 60, 6: 70, 7: 12, 8: 20, 9: 16})

    rows = []
    for n in a.nodes:
        cur = n.has_current
        rows.append([n.plant, n.product, n.d_month, n.n_customers,
                     n.cur_fg_t if cur else None, n.cur_consig_t if cur else None,
                     n.cur_wip_t if cur else None, n.cur_transit_t if cur else None,
                     n.current_total_t if cur else None, n.optimal_total_t,
                     n.cycles_month_current, n.cycles_month_optimal,
                     n.cycles_month_current and n.cycles_month_current * 12,
                     n.cycles_month_optimal and n.cycles_month_optimal * 12,
                     n.doi_current, n.doi_optimal, n.delta_t,
                     n.delta_t * n.value_usd_t if n.delta_t is not None else None, n.status])
    out["Rotacion"] = dict(
        headers=["Planta", "Producto", "Demanda t/mes", "Clientes", "PT planta (t)", "Consignación (t)",
                 "En proceso (t)", "En tránsito (t)", "Total actual (t)", "Total óptimo (t)",
                 "Ciclos/mes actual", "Ciclos/mes óptimo", "Ciclos/año actual", "Ciclos/año óptimo",
                 "Días inv. actual", "Días inv. óptimo", "Exceso (+) / falta (-) t", "Exceso USD", "Estado PT"],
        rows=rows, fmts=[None, None, T, D0, T, T, T, T, T, T, X2, X2, "0.0", "0.0", D0, D0, T, USD, None],
        status_col=18, widths={1: 16, 2: 24, 19: 18}, freeze="C2")

    out["Stock_Optimo"] = dict(
        headers=["Planta", "Producto", "Demanda t/mes", "Campaña cada (días)", "Proceso + QA (días)",
                 "Stock seguridad (t)", "Seguridad sin pooling (t)", "Stock de ciclo (t)", "En proceso (t)",
                 "En tránsito (t)", "Mínimo PT (t)", "Objetivo PT (t)", "Máximo PT (t)",
                 "Punto de reorden: lanzar campaña (t)", "Total óptimo (t)", "Total óptimo USD"],
        rows=[[n.plant, n.product, n.d_month, n.campaign_days, n.lead_prod_days, n.ss_t, n.ss_unpooled_t,
               n.cycle_t, n.wip_t, n.transit_t, n.min_t, n.target_t, n.max_t, n.rop_t, n.optimal_total_t,
               n.optimal_total_t * n.value_usd_t] for n in a.nodes if n.d_month],
        fmts=[None, None, T, D0, D0, T, T, T, T, T, T, T, T, T, T, USD], widths={1: 16, 2: 24}, freeze="C2")

    rows = []
    for pl in a.plants:
        diff = pl.conc_current_t - pl.conc_target_t if pl.conc_current_t is not None else None
        rows.append([pl.plant, pl.conc_d_day, pl.conc_d_day * p.dias_mes, pl.conc_lead_time, pl.conc_ss_t,
                     pl.conc_cycle_t, pl.conc_target_t, pl.conc_ss_t + 2 * pl.conc_cycle_t, pl.conc_current_t,
                     pl.conc_coverage_days, pl.conc_target_t / pl.conc_d_day if pl.conc_d_day else None, diff,
                     diff * pl.conc_value_usd_t if diff is not None else None])
    out["Concentrado"] = dict(
        headers=["Planta", "Consumo t/día", "Consumo t/mes", "Lead time (días)", "Stock seguridad (t)",
                 "Stock de ciclo (t)", "Objetivo (t)", "Máximo (t)", "Actual (t)", "Cobertura actual (días)",
                 "Cobertura objetivo (días)", "Exceso (+) / falta (-) t", "Exceso USD"],
        rows=rows, fmts=[None, "0.0", T, D0, T, T, T, T, T, D0, D0, T, USD], widths={1: 16})

    out["Clientes"] = dict(
        headers=["Cliente", "País", "Región", "Producto", "Planta", "ABC (volumen)", "XYZ (variabilidad)",
                 "Nivel servicio", "Demanda t/mes", "Desv. t/mes", "CV", "Meses historia", "Tránsito (días)",
                 "Desv. tránsito", "Ruta según", "Flete USD/t", "Despacho cada (días)", "Pedido firme (días)",
                 "Exposición (días)", "Seguridad propia (t)", "Ciclo (t)", "Tránsito (t)",
                 "% var. demanda", "% var. proceso", "% var. tránsito", "Costo anual USD"],
        rows=[[ln.customer, ln.country, ln.region, ln.product, ln.plant, ln.abc, ln.xyz, ln.service_level,
               ln.d_month, ln.sd_month, ln.cv, ln.months, ln.transit_days, ln.transit_sd, ln.lane_source,
               ln.freight_usd_t, ln.dispatch_days, ln.firm_days, ln.exposure_days, ln.ss_t, ln.cycle_t,
               ln.transit_stock_t, ln.share_demand, ln.share_process, ln.share_transit,
               a.line_annual_cost(ln)["total"]] for ln in a.lines],
        fmts=[None] * 7 + [PCT, T, T, PCT, D0, D0, D0, None, T, D0, D0, D0, T, T, T, PCT, PCT, PCT, USD],
        widths={1: 24, 4: 22, 5: 16}, freeze="B2")

    out["Reasignacion"] = dict(
        headers=["Cliente", "Producto", "t/mes", "Planta actual", "Planta propuesta", "Tránsito actual",
                 "Tránsito propuesto", "Costo anual actual", "Costo anual propuesto", "Ahorro anual", "Motivo"],
        rows=[[m.customer, m.product, m.t_month, m.from_plant, m.to_plant, m.transit_from, m.transit_to,
               m.cost_from, m.cost_to, m.saving, m.reason] for m in alloc.moves],
        fmts=[None, None, T, None, None, D0, D0, USD, USD, USD], widths={1: 24, 2: 22, 11: 70})
    return out


def load_table(alloc: AllocationResult) -> dict:
    return dict(
        headers=["Planta", "Capacidad t/mes", "Carga actual t/mes", "Utilización actual",
                 "Carga propuesta t/mes", "Utilización propuesta"],
        rows=[[pb.plant, pb.capacity_t_month, pb.load_t_month, pb.utilization, pp.load_t_month, pp.utilization]
              for pb, pp in zip(alloc.base.plants, alloc.proposed.plants)],
        fmts=[None, T, T, PCT, T, PCT])


def kpi_rows(a: Analysis) -> list[list]:
    t = a.totals()
    diff = (lambda x, y: x - y if x is not None and y is not None else None)
    return [
        ["Demanda (t/mes)", t["demanda_t_mes"], None, None],
        ["Ventas valorizadas (USD/mes)", t["ventas_usd_mes"], None, None],
        ["Inventario PT+proceso+tránsito (t)", t["inv_actual_t"], t["inv_optimo_t"],
         diff(t["inv_actual_t"], t["inv_optimo_t"])],
        ["Inventario producto (USD)", t["pt_actual_usd"], t["pt_optimo_usd"], diff(t["pt_actual_usd"], t["pt_optimo_usd"])],
        ["Inventario concentrado (USD)", t["conc_actual_usd"], t["conc_optimo_usd"],
         diff(t["conc_actual_usd"], t["conc_optimo_usd"])],
        ["Ciclos de rotación por mes (producto)", t["ciclos_mes_actual"], t["ciclos_mes_optimo"], None],
        ["Ciclos de rotación por año (producto)", t["ciclos_mes_actual"] and t["ciclos_mes_actual"] * 12,
         t["ciclos_mes_optimo"] and t["ciclos_mes_optimo"] * 12, None],
        ["Días de inventario (producto)", t["doi_actual"], t["doi_optimo"], None],
        ["Capital de trabajo liberable (USD)", t["capital_liberable_usd"], None, None],
        ["Ahorro anual de costo de capital (USD)", t["ahorro_anual_usd"], None, None],
    ]


PLANT_HEADERS = ["Planta", "Región", "Demanda t/mes", "Capacidad t/mes", "Utilización",
                 "Inventario actual USD", "Inventario óptimo USD", "Exceso USD",
                 "Ciclos/mes actual", "Ciclos/mes óptimo", "Días inv. actual", "Días inv. óptimo",
                 "Cobertura concentrado (días)"]


def plant_summary(a: Analysis) -> list[list]:
    agg = defaultdict(lambda: defaultdict(float))
    for n in a.nodes:
        g = agg[n.plant]
        g["d"] += n.d_month
        g["sales"] += n.d_month * n.value_usd_t
        g["cur"] += n.current_total_t * n.value_usd_t
        g["opt"] += n.optimal_total_t * n.value_usd_t
        g["has"] += 1 if n.has_current else 0
    out = []
    for p in a.plants:
        g = agg[p.plant]
        has = g["has"] > 0
        out.append([
            p.plant, p.region, g["d"], p.capacity_t_month, p.utilization,
            g["cur"] if has else None, g["opt"], (g["cur"] - g["opt"]) if has else None,
            g["sales"] / g["cur"] if has and g["cur"] else None, g["sales"] / g["opt"] if g["opt"] else None,
            g["cur"] / (g["sales"] / 30) if has and g["sales"] else None, g["opt"] / (g["sales"] / 30) if g["sales"] else None,
            p.conc_coverage_days,
        ])
    return out


def scenario(a: Analysis) -> dict[str, float]:
    cost = defaultdict(float)
    for ln in a.lines:
        for k, v in a.line_annual_cost(ln).items():
            cost[k] += v
    vol = sum(ln.d_month for ln in a.lines) or 1.0
    t = a.totals()
    return {
        "Flete anual (USD)": cost["flete"],
        "Conversión + logística concentrado (USD/año)": cost["conversion"],
        "Costo de capital del inventario óptimo (USD/año)": cost["inventario"],
        "Costo total anual (USD)": cost["total"],
        "Inventario óptimo PT+proceso+tránsito (t)": t["inv_optimo_t"],
        "Inventario óptimo total (USD)": t["inv_optimo_usd"],
        "Ciclos por mes óptimos": t["ciclos_mes_optimo"],
        "Tránsito medio ponderado (días)": sum(ln.transit_days * ln.d_month for ln in a.lines) / vol,
    }


def write_report(alloc: AllocationResult, recs: list[Recommendation], path) -> None:
    a = alloc.base
    p = a.ds.params
    wb = Workbook()

    # ------------------------------------------------ Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Análisis de rotación y stock óptimo - productos de molibdeno"
    ws["A1"].font = Font(bold=True, size=14)
    period = f"{a.months[0]} a {a.months[-1]}" if a.months else "sin historia"
    ws["A2"] = (f"Generado {dt.date.today():%Y-%m-%d} · historia de demanda {period} ({len(a.months)} meses) · "
                f"{len(a.lines)} líneas cliente-producto · {len(a.plants)} plantas")
    num = "#,##0.00"
    r = _table(ws, 4, ["Indicador", "Actual", "Óptimo", "Diferencia (actual - óptimo)"], kpi_rows(a),
               [None, num, num, num])
    ws.cell(r - 1, 1, "Asignación actual de clientes; capital liberable = inventario actual - óptimo "
                      "(solo donde hay stock informado).").font = Font(italic=True, size=9)

    plant_rows = plant_summary(a)
    ws.cell(r, 1, "Por planta").font = Font(bold=True, size=12)
    r0 = r + 1
    r = _table(ws, r0, PLANT_HEADERS, plant_rows,
               [None, None, T, T, PCT, USD, USD, USD, X2, X2, D0, D0, D0])
    if plant_rows:
        chart = BarChart()
        chart.title = "Ciclos de rotación por mes: actual vs óptimo"
        chart.y_axis.title = "ciclos/mes"
        data = Reference(ws, min_col=9, max_col=10, min_row=r0, max_row=r0 + len(plant_rows))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=r0 + 1, max_row=r0 + len(plant_rows)))
        chart.height, chart.width = 7.5, 16
        ws.add_chart(chart, f"A{r + 1}")
        r += 17

    ws.cell(r, 1, "Escenario: asignación actual vs propuesta").font = Font(bold=True, size=12)
    base, prop = scenario(alloc.base), scenario(alloc.proposed)
    _table(ws, r + 1, ["Indicador", "Asignación actual", "Asignación propuesta", "Diferencia"],
           [[k, base[k], prop[k], prop[k] - base[k] if None not in (base[k], prop[k]) else None] for k in base],
           [None, num, num, num])
    _widths(ws, {1: 44, 2: 18, 3: 18, 4: 22})

    # ------------------------------------------------ hojas de detalle
    tables = detail_tables(alloc, recs)
    for name, tb in tables.items():
        ws = wb.create_sheet(name)
        r = _table(ws, 1, tb["headers"], tb["rows"] or [["Sin datos"]], tb["fmts"],
                   status_col=tb.get("status_col"), fills=tb.get("fills"))
        _widths(ws, tb.get("widths"))
        if tb.get("freeze"):
            ws.freeze_panes = tb["freeze"]
        if tb.get("wrap"):
            for row in ws.iter_rows(min_row=2, min_col=tb["wrap"][0], max_col=tb["wrap"][1]):
                for c in row:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        if name == "Reasignacion":
            ws.cell(r, 1, f"Solo se proponen cambios con ahorro > {p.ahorro_minimo_reasignacion_usd:,.0f} USD/año, "
                          "respetando capacidades. El costo incluye flete, conversión, logística del concentrado y "
                          "costo de capital del stock.").font = Font(italic=True, size=9)
            lt = load_table(alloc)
            _table(ws, r + 2, lt["headers"], lt["rows"], lt["fmts"])

    # ------------------------------------------------ Supuestos y advertencias
    ws = wb.create_sheet("Supuestos")
    rows = [[f.name, getattr(p, f.name), PARAM_HELP.get(f.name, "")] for f in fields(p)]
    r = _table(ws, 1, ["Parámetro", "Valor", "Descripción"], rows)
    ws.cell(r, 1, "Advertencias de datos").font = Font(bold=True, size=12)
    _table(ws, r + 1, ["#", "Advertencia"], [[i, w] for i, w in enumerate(a.ds.warnings, 1)] or [["-", "Sin advertencias"]])
    _widths(ws, {1: 36, 2: 90, 3: 70})

    # ------------------------------------------------ Metodología
    ws = wb.create_sheet("Metodologia")
    for i, line in enumerate(METHOD.strip().splitlines(), start=1):
        ws.cell(i, 1, line)
    ws["A1"].font = Font(bold=True, size=12)
    ws.column_dimensions["A"].width = 130

    wb.save(path)


METHOD = """
Metodología
1. Demanda: media y desviación estándar mensual por cliente-producto (desde su primer mes con datos; meses sin despacho = 0).
2. Clasificación ABC por volumen del cliente (80% / 15% / 5%) y XYZ por coeficiente de variación (<=25%, <=50%, >50%).
3. Nivel de servicio: el del cliente si se informa; si no, el del segmento A/B/C (hoja Parametros). z = cuantil normal.
4. Exposición = intervalo de campaña + proceso + QA + tránsito - días de pedido firme del cliente.
5. Stock de seguridad por línea = z · sqrt(Exposición · σd² + d² · (σproceso² + σtránsito²)), con d y σd diarios.
6. Por planta-producto el stock de seguridad se agrega con pooling: sqrt(Σ SS²). La diferencia con la suma simple es el ahorro de pooling.
7. Stock de ciclo = d · max(intervalo de campaña, intervalo de despacho) / 2.  En proceso = d · tiempo de proceso.  En tránsito = d · días de tránsito.
8. Niveles de producto terminado: mínimo = seguridad, objetivo = seguridad + ciclo, máximo = seguridad + 2 · ciclo; punto de reorden (lanzar campaña) = seguridad + d · (proceso + QA).
9. Rotación (ciclos/mes) = demanda mensual / inventario total (PT + consignación + proceso + tránsito). Ciclos/año = x12. Días de inventario = inventario / demanda diaria. A nivel agregado se usa valor (USD) para sumar productos distintos.
10. Concentrado: mismo modelo con el consumo de cada planta (t producto x factor de concentrado), su lead time y la frecuencia de recepción.
11. Reasignación: heurística de regret con capacidad por planta y por producto; costo anual = flete + conversión + logística del concentrado + costo de capital del stock que exige la línea. Solo se propone un cambio si ahorra más que el umbral.
12. Limitaciones: la demanda se supone estacionaria e independiente entre clientes; no se modelan costos fijos de planta ni costos de setup de campaña (las propuestas de campañas más frecuentes indican el costo máximo de setup que se justifica).
"""
